import re
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from ..models import PresentationData, SlideData, SlideType, ThemeConfig, ArchetypeType, DensityTier, BulletStyle
from ..markdown_utils import parse_markdown_runs
from .text_fit import line_height_pt, wrap_markdown


def _get_font_name(theme: ThemeConfig) -> str:
    if theme.archetype == ArchetypeType.EDITORIAL:
        return "Georgia"
    elif theme.archetype == ArchetypeType.TERMINAL:
        return "Consolas"
    elif theme.archetype == ArchetypeType.NEUBRUTALISM:
        return "Arial Black"
    elif theme.archetype == ArchetypeType.MINIMAL:
        return "Arial"
    elif theme.archetype == ArchetypeType.CORPORATE:
        return "Calibri"
    return "Calibri"


def _get_title_font(theme: ThemeConfig) -> str:
    """Font for slide and card titles. Like the HTML, non-editorial/terminal archetypes take the theme's
    title font, so a serif theme (warm_academic) keeps serif titles on any archetype."""
    if theme.archetype not in (ArchetypeType.EDITORIAL, ArchetypeType.TERMINAL):
        first = theme.fonts.get("title_font", "").split(",")[0].strip(" '\"").lower()
        if first in ("georgia", "times new roman", "serif"):
            return "Georgia"
    return _get_font_name(theme)


def _add_formatted_runs(paragraph, text: str, default_font: str, default_size: Pt, default_color: RGBColor, is_title: bool = False, accent_color: Optional[RGBColor] = None):
    """Parses inline Markdown and appends formatted TextRun objects to a PPTX paragraph."""
    runs_data = parse_markdown_runs(text)
    if not runs_data:
        run = paragraph.add_run()
        run.text = ""
        return

    for item in runs_data:
        if getattr(item, "is_break", False):
            paragraph.add_line_break()
            continue

        chunk_text = item.text
        if not chunk_text:
            continue

        is_bold = item.is_bold
        is_italic = item.is_italic
        is_code = item.is_code
        is_strike = getattr(item, "is_strike", False)
        is_highlight = getattr(item, "is_highlight", False)
        is_sub = getattr(item, "is_sub", False)
        is_sup = getattr(item, "is_sup", False)
        is_arrow = getattr(item, "is_arrow", False)

        run = paragraph.add_run()
        run.text = chunk_text
        run.font.name = "Consolas" if is_code else default_font
        if is_sup or is_sub:
            run.font.size = Pt(max(default_size.pt * 0.72, 8))
        else:
            run.font.size = default_size

        if is_arrow:
            run.font.color.rgb = accent_color if accent_color is not None else default_color
            run.font.bold = True
        else:
            run.font.color.rgb = default_color
            run.font.bold = is_bold or is_title

        run.font.italic = is_italic

        if item.is_link and item.link_url:
            run.hyperlink.address = item.link_url
            run.font.underline = True

        rPr = run._r.get_or_add_rPr()
        if is_strike:
            rPr.set('strike', 'sngStrike')
        if is_sup:
            rPr.set('baseline', '30000')
        elif is_sub:
            rPr.set('baseline', '-25000')
        if is_highlight:
            try:
                rPr.append(parse_xml(f'<a:highlight {nsdecls("a")}><a:srgbClr val="FFD700"/></a:highlight>'))
            except Exception:
                pass


_BULLET_RE = re.compile(r"^[-*•]\s+(.*)$")


def _bullet_marker(bullet_style: BulletStyle) -> str:
    return "› " if bullet_style == BulletStyle.CHEVRON else ("– " if bullet_style == BulletStyle.EDITORIAL else "• ")


def _render_card_description_pptx(desc_box, desc_text: str, font_name: str, card_desc_size: Pt, secondary_rgb: RGBColor, accent_rgb: RGBColor, bullet_style: BulletStyle):
    """Renders card description, formatting bullet lines with styled bullet markers in accent color."""
    # ponytail: clean bullet formatting per bullet_style
    tf = desc_box.text_frame
    tf.word_wrap = True
    lines = desc_text.split("\n")
    has_bullet = any(_BULLET_RE.match(line.strip()) for line in lines)

    if not has_bullet:
        dp = tf.paragraphs[0]
        dp.line_spacing = 1.3
        _add_formatted_runs(dp, desc_text, font_name, card_desc_size, secondary_rgb, accent_color=accent_rgb)
        return

    bullet_marker = _bullet_marker(bullet_style)
    first = True
    for line in lines:
        s = line.strip()
        if not s:
            continue
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.line_spacing = 1.25
        p.space_after = Pt(3)
        m = _BULLET_RE.match(s)
        if m:
            bullet_run = p.add_run()
            bullet_run.text = bullet_marker
            bullet_run.font.name = font_name
            bullet_run.font.size = card_desc_size
            bullet_run.font.bold = True
            bullet_run.font.color.rgb = accent_rgb
            _add_formatted_runs(p, m.group(1).strip(), font_name, card_desc_size, secondary_rgb, accent_color=accent_rgb)
        else:
            _add_formatted_runs(p, s, font_name, card_desc_size, secondary_rgb, accent_color=accent_rgb)


def _card_description_height_pt(desc_text: str, font_name: str, size_pt: float, width_pt: float, bullet_style: BulletStyle) -> float:
    """Measured height of `_render_card_description_pptx` output (same paragraphs, spacing and markers)."""
    if not desc_text:
        return 0.0
    line_h = line_height_pt(font_name, size_pt)
    lines = desc_text.split("\n")
    if not any(_BULLET_RE.match(line.strip()) for line in lines):
        return wrap_markdown(desc_text, font_name, size_pt, width_pt).lines * line_h * 1.3
    total = 0.0
    for line in lines:
        s = line.strip()
        if not s:
            continue
        m = _BULLET_RE.match(s)
        wrapped = wrap_markdown(m.group(1).strip(), font_name, size_pt, width_pt, prefix=_bullet_marker(bullet_style)) if m \
            else wrap_markdown(s, font_name, size_pt, width_pt)
        total += wrapped.lines * line_h * 1.25 + 3
    return total


# Default text frame insets of a python-pptx text box: 0.1" left/right, 0.05" top/bottom.
_BOX_INSET_X = Inches(0.2)
_BOX_INSET_Y = Inches(0.1)
_TITLE_DESC_GAP = Inches(0.12)
_MIN_CARD_DESC_PT = 12

# Card text placement per archetype: (left offset, width reduction, title top offset, bottom padding).
_CARD_TEXT_GEOMETRY = {
    ArchetypeType.EDITORIAL: (Inches(0.2), Inches(0.3), Inches(0.55), Inches(0.0)),
    ArchetypeType.NEUBRUTALISM: (Inches(0.25), Inches(0.5), Inches(0.95), Inches(0.2)),
    ArchetypeType.MINIMAL: (Inches(0.0), Inches(0.0), Inches(0.75), Inches(0.0)),
    ArchetypeType.CORPORATE: (Inches(0.25), Inches(0.5), Inches(0.9), Inches(0.2)),
}
_DEFAULT_CARD_TEXT_GEOMETRY = (Inches(0.25), Inches(0.5), Inches(0.95), Inches(0.2))  # cards, terminal


@dataclass
class _CardTextFit:
    title_pt: int
    desc_pt: int
    title_heights_pt: List[float]
    fits: bool


def _fit_card_text(cards, title_prefix: str, title_font: str, body_font: str, base_title_pt: float, base_desc_pt: float,
                   box_w, avail_h, bullet_style: BulletStyle) -> _CardTextFit:
    """Largest common title/description size (from the density tier down to 12 pt) at which every card
    title fits without a mid-word break and title + description fit the card height."""
    width_pt = (box_w - _BOX_INSET_X) / 12700
    avail_pt = (avail_h - 2 * _BOX_INSET_Y - _TITLE_DESC_GAP) / 12700
    desc_pt = int(base_desc_pt)
    while True:
        title_pt = max(desc_pt, round(desc_pt * base_title_pt / base_desc_pt))
        title_line_h = line_height_pt(title_font, title_pt)
        heights, fits = [], True
        for c in cards:
            title = wrap_markdown(f"{title_prefix}{c.title}", title_font, title_pt, width_pt, bold=True)
            title_h = title.lines * title_line_h
            heights.append(title_h)
            desc_h = _card_description_height_pt(c.description, body_font, desc_pt, width_pt, bullet_style)
            if title.broke_word or title_h + desc_h > avail_pt:
                fits = False
        if fits or desc_pt <= _MIN_CARD_DESC_PT:
            return _CardTextFit(title_pt, desc_pt, heights, fits)
        desc_pt -= 1


def _add_card_text(slide, card, left, top, width, bottom, title_h_pt: float, fit: _CardTextFit, title_prefix: str,
                   title_font: str, body_font: str, primary_rgb, secondary_rgb, accent_rgb, bullet_style: BulletStyle):
    """Card title box sized to its measured height, description stacked directly below it."""
    title_h = Pt(title_h_pt) + _BOX_INSET_Y
    title_box = slide.shapes.add_textbox(left, top, width, title_h)
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    _add_formatted_runs(tf_t.paragraphs[0], f"{title_prefix}{card.title}", title_font, Pt(fit.title_pt), primary_rgb, is_title=True, accent_color=accent_rgb)

    desc_top = top + title_h + _TITLE_DESC_GAP
    desc_box = slide.shapes.add_textbox(left, desc_top, width, max(bottom - desc_top, _BOX_INSET_Y))
    _render_card_description_pptx(desc_box, card.description, body_font, Pt(fit.desc_pt), secondary_rgb, accent_rgb, bullet_style)


def _add_picture_fit(slide, img_path: Path, left, top, max_width, max_height):
    """Adds a picture to the slide fitting within max_width and max_height while preserving aspect ratio."""
    try:
        with Image.open(img_path) as im:
            orig_w, orig_h = im.size
        if orig_h == 0 or orig_w == 0:
            return slide.shapes.add_picture(str(img_path), left, top, width=max_width)

        aspect = orig_w / orig_h
        box_aspect = max_width / max_height
        if aspect > box_aspect:
            render_w = int(max_width)
            render_h = int(max_width / aspect)
            render_left = int(left)
            render_top = int(top + (max_height - render_h) / 2)
        else:
            render_h = int(max_height)
            render_w = int(max_height * aspect)
            render_top = int(top)
            render_left = int(left + (max_width - render_w) / 2)
        return slide.shapes.add_picture(str(img_path), render_left, render_top, width=render_w, height=render_h)
    except Exception as e:
        print(f"[PPTX] Warning: Could not add picture {img_path}: {e}")
        return None


def render_pptx_presentation(pres: PresentationData, theme: ThemeConfig, dist_dir: Path) -> Path:
    dist_dir.mkdir(parents=True, exist_ok=True)
    out_file = dist_dir / "presentation.pptx"

    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    bg_rgb = RGBColor(*theme.hex_to_rgb(theme.colors['background']))
    card_bg_rgb = RGBColor(*theme.hex_to_rgb(theme.colors['card_bg']))
    card_border_rgb = RGBColor(*theme.hex_to_rgb(theme.colors.get('card_border', '#E2E8F0')))
    primary_text_rgb = RGBColor(*theme.hex_to_rgb(theme.colors['primary_text']))
    secondary_text_rgb = RGBColor(*theme.hex_to_rgb(theme.colors['secondary_text']))
    accent_rgb = RGBColor(*theme.hex_to_rgb(theme.colors['accent']))
    badge_bg_rgb = RGBColor(*theme.hex_to_rgb(theme.colors.get('badge_bg', '#DBEAFE')))
    badge_text_rgb = RGBColor(*theme.hex_to_rgb(theme.colors.get('badge_text', '#1E40AF')))

    total_slides = len(pres.slides)

    for slide_data in pres.slides:
        slide = prs.slides.add_slide(blank_layout)

        # 1. Slide background shape
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = bg_rgb
        bg_shape.line.fill.background()  # no line

        # 2. Add Audio shape if present
        if _has_audio(slide_data):
            _add_audio_widget(slide, slide_data.audio_path, accent_rgb, badge_bg_rgb, badge_text_rgb, theme)

        # 3. Slide counter in bottom right
        counter_box = slide.shapes.add_textbox(Inches(11.5), Inches(6.8), Inches(1.2), Inches(0.4))
        cp = counter_box.text_frame.paragraphs[0]
        cp.text = f"{slide_data.index} / {total_slides}"
        cp.font.name = _get_font_name(theme)
        cp.font.size = Pt(11)
        cp.font.color.rgb = secondary_text_rgb
        cp.alignment = PP_ALIGN.RIGHT

        # 4. Render according to slide type and archetype
        deck_title = pres.title
        if slide_data.slide_type == SlideType.TITLE:
            _render_pptx_title_slide(
                slide, slide_data, theme,
                primary_text_rgb, secondary_text_rgb, accent_rgb, badge_bg_rgb, badge_text_rgb,
                deck_title=deck_title
            )
        elif slide_data.slide_type == SlideType.DEFINITION:
            _render_pptx_definition_slide(
                slide, slide_data, theme,
                card_bg_rgb, card_border_rgb, primary_text_rgb, secondary_text_rgb, accent_rgb,
                deck_title=deck_title
            )
        elif slide_data.slide_type == SlideType.CARDS:
            _render_pptx_cards_slide(
                slide, slide_data, theme,
                card_bg_rgb, card_border_rgb, primary_text_rgb, secondary_text_rgb, accent_rgb, badge_bg_rgb, badge_text_rgb,
                deck_title=deck_title
            )
        elif slide_data.slide_type == SlideType.CODE:
            _render_pptx_code_slide(
                slide, slide_data, theme,
                primary_text_rgb, secondary_text_rgb, accent_rgb,
                deck_title=deck_title
            )
        elif slide_data.slide_type == SlideType.TABLE:
            _render_pptx_table_slide(
                slide, slide_data, theme,
                card_bg_rgb, card_border_rgb, primary_text_rgb, secondary_text_rgb, accent_rgb,
                deck_title=deck_title
            )
        elif slide_data.slide_type == SlideType.IMAGE:
            _render_pptx_image_slide(
                slide, slide_data, theme,
                accent_rgb, primary_text_rgb, secondary_text_rgb,
                deck_title=deck_title
            )
        else:
            _render_pptx_standard_slide(
                slide, slide_data, theme,
                card_bg_rgb, card_border_rgb, primary_text_rgb, secondary_text_rgb, accent_rgb,
                deck_title=deck_title
            )

    prs.save(out_file)
    return out_file


def _add_audio_widget(slide, audio_path: Path, accent_rgb, badge_bg_rgb, badge_text_rgb, theme: ThemeConfig):
    font_name = _get_font_name(theme)
    # If neubrutalism, hard shadow for audio badge
    if theme.archetype == ArchetypeType.NEUBRUTALISM:
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.24), Inches(0.59), Inches(2.2), Inches(0.6))
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor(0, 0, 0)
        sh.line.fill.background()

    shape_type = MSO_SHAPE.RECTANGLE if theme.archetype in (ArchetypeType.NEUBRUTALISM, ArchetypeType.TERMINAL, ArchetypeType.CORPORATE) else MSO_SHAPE.ROUNDED_RECTANGLE
    badge = slide.shapes.add_shape(shape_type, Inches(10.2), Inches(0.55), Inches(2.2), Inches(0.6))
    badge.fill.solid()
    badge.fill.fore_color.rgb = badge_bg_rgb
    
    if theme.archetype == ArchetypeType.NEUBRUTALISM:
        badge.line.color.rgb = RGBColor(0, 0, 0)
        badge.line.width = Pt(2)
    else:
        badge.line.color.rgb = accent_rgb
        badge.line.width = Pt(1)
    
    tf = badge.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "  🔊 Audio Narration"
    p.font.name = font_name
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = badge_text_rgb
    p.alignment = PP_ALIGN.LEFT

    # Add native movie/audio shape (loudspeaker icon that plays on click)
    try:
        slide.shapes.add_movie(
            movie_file=str(audio_path),
            left=Inches(11.8),
            top=Inches(0.58),
            width=Inches(0.52),
            height=Inches(0.52),
            mime_type="audio/mp4"
        )
    except Exception as e:
        print(f"[PPTX] Warning: Could not embed audio {audio_path.name}: {e}")


def _render_pptx_title_slide(slide, slide_data: SlideData, theme: ThemeConfig,
                             primary_rgb, secondary_rgb, accent_rgb, badge_bg_rgb, badge_text_rgb,
                             deck_title: str = ""):
    font_name = _get_font_name(theme)
    arch = theme.archetype
    badge_label = deck_title.upper() if deck_title else "PREZENTÁCIÓ"
    # Title slides are always classified SPACIOUS (models.calculate_density_tier).
    title_size = Pt(58)
    author_size = Pt(28)

    # 1. Category badge
    if arch == ArchetypeType.EDITORIAL:
        # Subtle non-enclosed text with fine underline
        cat_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(4.5), Inches(0.45))
        p = cat_box.text_frame.paragraphs[0]
        p.text = badge_label
        p.font.name = font_name
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = accent_rgb
    elif arch == ArchetypeType.NEUBRUTALISM:
        # Shadow + thick border
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.04), Inches(1.84), Inches(3.4), Inches(0.45))
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor(0, 0, 0)
        sh.line.fill.background()

        cat_badge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.8), Inches(3.4), Inches(0.45))
        cat_badge.fill.solid()
        cat_badge.fill.fore_color.rgb = badge_bg_rgb
        cat_badge.line.color.rgb = RGBColor(0, 0, 0)
        cat_badge.line.width = Pt(2)
        p = cat_badge.text_frame.paragraphs[0]
        p.text = badge_label
        p.font.name = font_name
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = badge_text_rgb
        p.alignment = PP_ALIGN.CENTER
    else:
        shape_type = MSO_SHAPE.RECTANGLE if arch in (ArchetypeType.TERMINAL, ArchetypeType.CORPORATE, ArchetypeType.MINIMAL) else MSO_SHAPE.ROUNDED_RECTANGLE
        cat_badge = slide.shapes.add_shape(shape_type, Inches(1.0), Inches(1.8), Inches(3.4), Inches(0.45))
        cat_badge.fill.solid()
        cat_badge.fill.fore_color.rgb = badge_bg_rgb
        cat_badge.line.fill.background()
        p = cat_badge.text_frame.paragraphs[0]
        prefix = "[INIT] " if arch == ArchetypeType.TERMINAL else ""
        p.text = f"{prefix}{badge_label}"
        p.font.name = font_name
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = badge_text_rgb
        p.alignment = PP_ALIGN.CENTER

    # 2. Main Title
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(11.0), Inches(2.3))
    tf = title_box.text_frame
    tf.word_wrap = True
    tp = tf.paragraphs[0]
    _add_formatted_runs(tp, slide_data.title, _get_title_font(theme), title_size, primary_rgb, is_title=True, accent_color=accent_rgb)

    # 3. Accent decorative line
    if arch != ArchetypeType.MINIMAL:
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(5.0), Inches(1.2), Inches(0.06))
        line.fill.solid()
        line.fill.fore_color.rgb = accent_rgb
        line.line.fill.background()

    # 4. Author / Subtitle
    if slide_data.author or slide_data.subtitle:
        author_text = slide_data.author or slide_data.subtitle
        author_box = slide.shapes.add_textbox(Inches(1.0), Inches(5.2), Inches(10.0), Inches(0.8))
        ap = author_box.text_frame.paragraphs[0]
        _add_formatted_runs(ap, author_text, font_name, author_size, secondary_rgb, accent_color=accent_rgb)


def _render_pptx_definition_slide(slide, slide_data: SlideData, theme: ThemeConfig,
                                  card_bg_rgb, card_border_rgb, primary_rgb, secondary_rgb, accent_rgb,
                                  deck_title: str = ""):
    cat_label = deck_title.upper() if deck_title else "DEFINÍCIÓ"
    _add_header(slide, cat_label, slide_data, accent_rgb, primary_rgb, theme)
    font_name = _get_font_name(theme)
    arch = theme.archetype
    first_para = slide_data.quote or (slide_data.paragraphs[0] if slide_data.paragraphs else "")
    tier = slide_data.density_tier

    # Definition slides are never COMPACT (models.calculate_density_tier).
    def_font_size = Pt(30) if tier == DensityTier.SPACIOUS else Pt(24)
    def_spacing = 1.45 if tier == DensityTier.SPACIOUS else 1.35

    if arch == ArchetypeType.EDITORIAL:
        # Elegant hairline top & bottom dividers, serif italic text, transparent card
        top_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.4), Inches(11.333), Inches(0.015))
        top_line.fill.solid()
        top_line.fill.fore_color.rgb = card_border_rgb
        top_line.line.fill.background()

        bottom_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(6.2), Inches(11.333), Inches(0.015))
        bottom_line.fill.solid()
        bottom_line.fill.fore_color.rgb = card_border_rgb
        bottom_line.line.fill.background()

        text_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.8), Inches(11.333), Inches(3.2))
        tf = text_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.line_spacing = def_spacing
        _add_formatted_runs(p, first_para, font_name, def_font_size, primary_rgb, accent_color=accent_rgb)

    elif arch == ArchetypeType.NEUBRUTALISM:
        # Black hard shadow offset rectangle
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.07), Inches(2.27), Inches(11.333), Inches(4.3))
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor(0, 0, 0)
        sh.line.fill.background()

        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.2), Inches(11.333), Inches(4.3))
        card.fill.solid()
        card.fill.fore_color.rgb = card_bg_rgb
        card.line.color.rgb = RGBColor(0, 0, 0)
        card.line.width = Pt(3)

        # Thick black left bar
        left_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.2), Inches(0.2), Inches(4.3))
        left_bar.fill.solid()
        left_bar.fill.fore_color.rgb = RGBColor(0, 0, 0)
        left_bar.line.fill.background()

        text_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.6), Inches(10.3), Inches(3.5))
        tf = text_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.line_spacing = def_spacing
        _add_formatted_runs(p, first_para, font_name, def_font_size, primary_rgb, accent_color=accent_rgb)

    elif arch == ArchetypeType.MINIMAL:
        # Transparent card, bold left vertical bar
        left_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.2), Inches(0.08), Inches(4.3))
        left_bar.fill.solid()
        left_bar.fill.fore_color.rgb = accent_rgb
        left_bar.line.fill.background()

        text_box = slide.shapes.add_textbox(Inches(1.4), Inches(2.4), Inches(10.9), Inches(3.8))
        tf = text_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.line_spacing = def_spacing
        _add_formatted_runs(p, first_para, font_name, def_font_size, primary_rgb, accent_color=accent_rgb)

    elif arch == ArchetypeType.CORPORATE:
        # Sharp rectangle + top accent banner strip
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.2), Inches(11.333), Inches(4.3))
        card.fill.solid()
        card.fill.fore_color.rgb = card_bg_rgb
        card.line.color.rgb = card_border_rgb
        card.line.width = Pt(1)

        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.2), Inches(11.333), Inches(0.08))
        banner.fill.solid()
        banner.fill.fore_color.rgb = accent_rgb
        banner.line.fill.background()

        text_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.6), Inches(10.3), Inches(3.5))
        tf = text_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.line_spacing = def_spacing
        _add_formatted_runs(p, first_para, font_name, def_font_size, primary_rgb, accent_color=accent_rgb)

    else:
        # Terminal & Cards (modern cards)
        shape_type = MSO_SHAPE.RECTANGLE if arch == ArchetypeType.TERMINAL else MSO_SHAPE.ROUNDED_RECTANGLE
        card = slide.shapes.add_shape(shape_type, Inches(1.0), Inches(2.2), Inches(11.333), Inches(4.3))
        card.fill.solid()
        card.fill.fore_color.rgb = card_bg_rgb
        card.line.color.rgb = card_border_rgb
        card.line.width = Pt(1)

        left_bar = slide.shapes.add_shape(shape_type, Inches(1.0), Inches(2.2), Inches(0.12), Inches(4.3))
        left_bar.fill.solid()
        left_bar.fill.fore_color.rgb = accent_rgb
        left_bar.line.fill.background()

        text_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.6), Inches(10.3), Inches(3.5))
        tf = text_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.line_spacing = def_spacing
        _add_formatted_runs(p, first_para, font_name, def_font_size, primary_rgb, accent_color=accent_rgb)


def _render_pptx_cards_slide(slide, slide_data: SlideData, theme: ThemeConfig,
                             card_bg_rgb, card_border_rgb, primary_rgb, secondary_rgb, accent_rgb, badge_bg_rgb, badge_text_rgb,
                             deck_title: str = ""):
    cat_label = deck_title.upper() if deck_title else "ÁTTEKINTÉS"
    _add_header(slide, cat_label, slide_data, accent_rgb, primary_rgb, theme)
    font_name = _get_font_name(theme)
    arch = theme.archetype

    has_intro = bool(slide_data.paragraphs)
    has_footnote = len(slide_data.paragraphs) > 1
    top = Inches(2.9) if has_intro else Inches(2.2)
    height = Inches(3.7) if has_intro else Inches(4.3)
    footnote_h = Inches(0.55)
    if has_footnote:
        height -= footnote_h
    tier = slide_data.density_tier

    card_title_size = Pt(28) if tier == DensityTier.SPACIOUS else (Pt(23) if tier == DensityTier.BALANCED else Pt(19))
    card_desc_size = Pt(20) if tier == DensityTier.SPACIOUS else (Pt(17) if tier == DensityTier.BALANCED else Pt(14))
    intro_size = Pt(20) if tier == DensityTier.SPACIOUS else (Pt(17) if tier == DensityTier.BALANCED else Pt(14))
    footnote_size = Pt(16) if tier == DensityTier.SPACIOUS else (Pt(14) if tier == DensityTier.BALANCED else Pt(12))

    if has_intro:
        intro_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.1), Inches(11.333), Inches(0.65))
        itf = intro_box.text_frame
        itf.word_wrap = True
        ip = itf.paragraphs[0]
        _add_formatted_runs(ip, slide_data.paragraphs[0], font_name, intro_size, secondary_rgb, accent_color=accent_rgb)

    if has_footnote:
        # Mirrors the HTML `.slide-footnote` (secondary color, italic) below the cards grid.
        fn_box = slide.shapes.add_textbox(Inches(1.0), top + height + Inches(0.1), Inches(11.333), footnote_h - Inches(0.1))
        ftf = fn_box.text_frame
        ftf.word_wrap = True
        fp = ftf.paragraphs[0]
        _add_formatted_runs(fp, slide_data.paragraphs[1], font_name, footnote_size, secondary_rgb, accent_color=accent_rgb)
        for run in fp.runs:
            run.font.italic = True

    num_cards = max(1, min(4, len(slide_data.cards)))
    if num_cards == 1:
        card_w = Inches(11.333)
        gap = Inches(0)
    elif num_cards == 2:
        card_w = Inches(5.45)
        gap = Inches(0.43)
    elif num_cards == 3:
        card_w = Inches(3.5)
        gap = Inches(0.41)
    else:  # 4 cards
        card_w = Inches(2.6)
        gap = Inches(0.31)

    text_dx, text_dw, title_dy, bottom_pad = _CARD_TEXT_GEOMETRY.get(arch, _DEFAULT_CARD_TEXT_GEOMETRY)
    title_prefix = "> " if arch == ArchetypeType.TERMINAL else ""
    cards = slide_data.cards[:4]
    fit = _fit_card_text(cards, title_prefix, _get_title_font(theme), font_name, card_title_size.pt, card_desc_size.pt,
                         card_w - text_dw, height - title_dy - bottom_pad, theme.bullet_style)
    if not fit.fits:
        print(f"[PPTX] Warning: slide {slide_data.index}: card text does not fit even at {fit.desc_pt} pt; shorten it.")

    for i, c in enumerate(cards):
        left = Inches(1.0) + i * (card_w + gap)

        if arch == ArchetypeType.EDITORIAL:
            # Transparent cards, left vertical hairline accent
            left_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top + Inches(0.1), Inches(0.03), height - Inches(0.2))
            left_line.fill.solid()
            left_line.fill.fore_color.rgb = accent_rgb
            left_line.line.fill.background()

            # Elegant Roman / numeric index
            idx_box = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.1), Inches(1.5), Inches(0.4))
            ip = idx_box.text_frame.paragraphs[0]
            ip.text = f"{c.card_index}."
            ip.font.name = font_name
            ip.font.size = Pt(14)
            ip.font.bold = True
            ip.font.color.rgb = accent_rgb

        elif arch == ArchetypeType.NEUBRUTALISM:
            # Hard black shadow shape
            sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + Inches(0.07), top + Inches(0.07), card_w, height)
            sh.fill.solid()
            sh.fill.fore_color.rgb = RGBColor(0, 0, 0)
            sh.line.fill.background()

            # Card shape on top
            card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, card_w, height)
            card.fill.solid()
            card.fill.fore_color.rgb = card_bg_rgb
            card.line.color.rgb = RGBColor(0, 0, 0)
            card.line.width = Pt(2.5)

            # Badge shadow + badge
            b_sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + Inches(0.44), top + Inches(0.44), Inches(0.55), Inches(0.45))
            b_sh.fill.solid()
            b_sh.fill.fore_color.rgb = RGBColor(0, 0, 0)
            b_sh.line.fill.background()

            badge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + Inches(0.4), top + Inches(0.4), Inches(0.55), Inches(0.45))
            badge.fill.solid()
            badge.fill.fore_color.rgb = accent_rgb
            badge.line.color.rgb = RGBColor(0, 0, 0)
            badge.line.width = Pt(2)
            bp = badge.text_frame.paragraphs[0]
            bp.text = f"0{c.card_index}" if c.card_index < 10 else str(c.card_index)
            bp.font.name = font_name
            bp.font.size = Pt(12)
            bp.font.bold = True
            bp.font.color.rgb = RGBColor(0, 0, 0)
            bp.alignment = PP_ALIGN.CENTER

        elif arch == ArchetypeType.MINIMAL:
            # Transparent cards, oversized numeric badge above title
            idx_box = slide.shapes.add_textbox(left, top, card_w, Inches(0.7))
            ip = idx_box.text_frame.paragraphs[0]
            ip.text = f"0{c.card_index}" if c.card_index < 10 else str(c.card_index)
            ip.font.name = font_name
            ip.font.size = Pt(36)
            ip.font.bold = True
            ip.font.color.rgb = accent_rgb

        elif arch == ArchetypeType.CORPORATE:
            # Sharp rectangle + top accent banner
            card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, card_w, height)
            card.fill.solid()
            card.fill.fore_color.rgb = card_bg_rgb
            card.line.color.rgb = card_border_rgb
            card.line.width = Pt(1)

            banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, card_w, Inches(0.08))
            banner.fill.solid()
            banner.fill.fore_color.rgb = accent_rgb
            banner.line.fill.background()

            badge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + Inches(0.35), top + Inches(0.35), Inches(0.55), Inches(0.4))
            badge.fill.solid()
            badge.fill.fore_color.rgb = accent_rgb
            badge.line.fill.background()
            bp = badge.text_frame.paragraphs[0]
            bp.text = f"0{c.card_index}" if c.card_index < 10 else str(c.card_index)
            bp.font.name = font_name
            bp.font.size = Pt(11)
            bp.font.bold = True
            bp.font.color.rgb = RGBColor(255, 255, 255)
            bp.alignment = PP_ALIGN.CENTER

        else:
            # Terminal & Modern Cards
            shape_type = MSO_SHAPE.RECTANGLE if arch == ArchetypeType.TERMINAL else MSO_SHAPE.ROUNDED_RECTANGLE
            card = slide.shapes.add_shape(shape_type, left, top, card_w, height)
            card.fill.solid()
            card.fill.fore_color.rgb = card_bg_rgb
            card.line.color.rgb = card_border_rgb
            card.line.width = Pt(1)

            badge = slide.shapes.add_shape(shape_type, left + Inches(0.4), top + Inches(0.4), Inches(0.55), Inches(0.45))
            badge.fill.solid()
            badge.fill.fore_color.rgb = badge_bg_rgb
            badge.line.fill.background()
            bp = badge.text_frame.paragraphs[0]
            bp.text = f"0{c.card_index}" if c.card_index < 10 else str(c.card_index)
            bp.font.name = font_name
            bp.font.size = Pt(12)
            bp.font.bold = True
            bp.font.color.rgb = badge_text_rgb
            bp.alignment = PP_ALIGN.CENTER

        _add_card_text(slide, c, left + text_dx, top + title_dy, card_w - text_dw, top + height - bottom_pad,
                       fit.title_heights_pt[i], fit, title_prefix, _get_title_font(theme), font_name, primary_rgb, secondary_rgb, accent_rgb,
                       theme.bullet_style)


def _render_pptx_code_slide(slide, slide_data: SlideData, theme: ThemeConfig,
                            primary_rgb, secondary_rgb, accent_rgb,
                            deck_title: str = ""):
    _add_header(slide, "KÓDRÉSZLET", slide_data, accent_rgb, primary_rgb, theme)
    lang = (slide_data.code_snippet.get("language") if slide_data.code_snippet else "code") or "code"
    code_text = (slide_data.code_snippet.get("code") if slide_data.code_snippet else "") or ""

    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.2), Inches(11.333), Inches(4.3))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(13, 17, 23)
    box.line.color.rgb = accent_rgb
    box.line.width = Pt(1)

    lbl = slide.shapes.add_textbox(Inches(1.2), Inches(2.3), Inches(4.0), Inches(0.4))
    lp = lbl.text_frame.paragraphs[0]
    lp.text = f"// LANGUAGE: {lang.upper()}"
    lp.font.name = "Consolas"
    lp.font.size = Pt(11)
    lp.font.bold = True
    lp.font.color.rgb = accent_rgb

    code_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.8), Inches(10.9), Inches(3.5))
    tf = code_box.text_frame
    tf.word_wrap = True
    lines = code_text.splitlines()[:14]
    for idx, line in enumerate(lines):
        p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
        p.text = line
        p.font.name = "Consolas"
        p.font.size = Pt(13)
        p.font.color.rgb = RGBColor(230, 237, 243)


def _render_pptx_table_slide(slide, slide_data: SlideData, theme: ThemeConfig,
                             card_bg_rgb, card_border_rgb, primary_rgb, secondary_rgb, accent_rgb,
                             deck_title: str = ""):
    cat_label = deck_title.upper() if deck_title else "TÁBLÁZAT"
    _add_header(slide, cat_label, slide_data, accent_rgb, primary_rgb, theme)
    font_name = _get_font_name(theme)
    arch = theme.archetype
    tier = slide_data.density_tier

    header_size = Pt(18) if tier == DensityTier.SPACIOUS else (Pt(16) if tier == DensityTier.BALANCED else Pt(13))
    cell_size = Pt(16) if tier == DensityTier.SPACIOUS else (Pt(14) if tier == DensityTier.BALANCED else Pt(12))
    intro_size = Pt(20) if tier == DensityTier.SPACIOUS else (Pt(17) if tier == DensityTier.BALANCED else Pt(14))

    table_data = slide_data.table
    if not table_data or not table_data.headers:
        return

    num_rows = len(table_data.rows) + 1
    num_cols = len(table_data.headers)

    top = Inches(2.2)
    left = Inches(1.0)
    width = Inches(11.333)

    if slide_data.paragraphs:
        intro_box = slide.shapes.add_textbox(left, top, width, Inches(0.55))
        tf_i = intro_box.text_frame
        tf_i.word_wrap = True
        ip = tf_i.paragraphs[0]
        _add_formatted_runs(ip, slide_data.paragraphs[0], font_name, intro_size, secondary_rgb, accent_color=accent_rgb)
        top = Inches(2.8)

    table_height = min(Inches(0.55 * num_rows), Inches(4.3))
    table_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, table_height)
    table = table_shape.table

    col_w = int(width / num_cols)
    for col in table.columns:
        col.width = col_w

    align_map = {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT
    }

    # Header Row
    for c_idx, h_text in enumerate(table_data.headers):
        cell = table.cell(0, c_idx)
        cell.fill.solid()
        if arch == ArchetypeType.EDITORIAL:
            cell.fill.fore_color.rgb = card_bg_rgb
            h_color = primary_rgb
        elif arch == ArchetypeType.TERMINAL:
            cell.fill.fore_color.rgb = RGBColor(20, 30, 20)
            h_color = accent_rgb
        else:
            cell.fill.fore_color.rgb = accent_rgb
            h_color = RGBColor(255, 255, 255)

        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        _add_formatted_runs(p, h_text, font_name, header_size, h_color, is_title=True, accent_color=accent_rgb)
        col_align = table_data.alignments[c_idx] if c_idx < len(table_data.alignments) else "left"
        p.alignment = align_map.get(col_align, PP_ALIGN.LEFT)

    # Body Rows
    for r_idx, row in enumerate(table_data.rows):
        is_even = (r_idx % 2 == 1)
        for c_idx, cell_text in enumerate(row):
            if c_idx >= num_cols:
                break
            cell = table.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            if is_even:
                cell.fill.fore_color.rgb = card_bg_rgb
            else:
                cell.fill.fore_color.rgb = RGBColor(
                    min(255, card_bg_rgb[0] + 8),
                    min(255, card_bg_rgb[1] + 8),
                    min(255, card_bg_rgb[2] + 8)
                ) if arch != ArchetypeType.NEUBRUTALISM else card_bg_rgb

            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            _add_formatted_runs(p, cell_text, font_name, cell_size, primary_rgb, accent_color=accent_rgb)
            col_align = table_data.alignments[c_idx] if c_idx < len(table_data.alignments) else "left"
            p.alignment = align_map.get(col_align, PP_ALIGN.LEFT)


def _render_pptx_image_slide(slide, slide_data: SlideData, theme: ThemeConfig,
                             accent_rgb, primary_rgb, secondary_rgb,
                             deck_title: str = ""):
    cat_label = deck_title.upper() if deck_title else "ILLUSZTRÁCIÓ"
    _add_header(slide, cat_label, slide_data, accent_rgb, primary_rgb, theme)
    font_name = _get_font_name(theme)
    tier = slide_data.density_tier
    intro_size = Pt(20) if tier == DensityTier.SPACIOUS else (Pt(17) if tier == DensityTier.BALANCED else Pt(14))

    top = Inches(2.2)
    max_h = Inches(4.5)

    if slide_data.paragraphs:
        intro_box = slide.shapes.add_textbox(Inches(1.0), top, Inches(11.333), Inches(0.5))
        tf_i = intro_box.text_frame
        tf_i.word_wrap = True
        ip = tf_i.paragraphs[0]
        ip.alignment = PP_ALIGN.CENTER
        _add_formatted_runs(ip, slide_data.paragraphs[0], font_name, intro_size, secondary_rgb, accent_color=accent_rgb)
        top = Inches(2.8)
        max_h = Inches(3.9)

    if slide_data.images:
        img_info = slide_data.images[0]
        if img_info.resolved_path and img_info.resolved_path.exists():
            _add_picture_fit(slide, img_info.resolved_path, Inches(1.5), top, Inches(10.333), max_h)
            if img_info.alt:
                cap_box = slide.shapes.add_textbox(Inches(1.0), Inches(6.8), Inches(11.333), Inches(0.4))
                cp = cap_box.text_frame.paragraphs[0]
                cp.alignment = PP_ALIGN.CENTER
                _add_formatted_runs(cp, img_info.alt, font_name, Pt(12), secondary_rgb, accent_color=accent_rgb)


def _render_pptx_standard_slide(slide, slide_data: SlideData, theme: ThemeConfig,
                                card_bg_rgb, card_border_rgb, primary_rgb, secondary_rgb, accent_rgb,
                                deck_title: str = ""):
    cat_label = deck_title.upper() if deck_title else "ÁTTEKINTÉS"
    _add_header(slide, cat_label, slide_data, accent_rgb, primary_rgb, theme)
    font_name = _get_font_name(theme)
    tier = slide_data.density_tier

    body_size = Pt(28) if tier == DensityTier.SPACIOUS else (Pt(22) if tier == DensityTier.BALANCED else Pt(18))
    para_space = Pt(20) if tier == DensityTier.SPACIOUS else (Pt(14) if tier == DensityTier.BALANCED else Pt(8))
    bullet_space = Pt(16) if tier == DensityTier.SPACIOUS else (Pt(10) if tier == DensityTier.BALANCED else Pt(6))

    has_image = bool(slide_data.images and slide_data.images[0].resolved_path and slide_data.images[0].resolved_path.exists())
    text_width = Inches(6.0) if has_image else Inches(11.333)

    tb = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), text_width, Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for p_text in slide_data.paragraphs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = para_space
        _add_formatted_runs(p, p_text, font_name, body_size, primary_rgb, accent_color=accent_rgb)

    for item_text in slide_data.list_items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = bullet_space
        p.level = 0
        bullet_run = p.add_run()
        if item_text.startswith("[x] "):
            bullet_run.text = "☑ "
            bullet_run.font.color.rgb = accent_rgb
            clean_item = item_text[4:]
        elif item_text.startswith("[ ] "):
            bullet_run.text = "☐ "
            bullet_run.font.color.rgb = secondary_rgb
            clean_item = item_text[4:]
        else:
            bullet_marker = "› " if theme.bullet_style == BulletStyle.CHEVRON else ("– " if theme.bullet_style == BulletStyle.EDITORIAL else "• ")
            bullet_run.text = bullet_marker
            bullet_run.font.color.rgb = accent_rgb
            clean_item = item_text
        bullet_run.font.name = font_name
        bullet_run.font.size = body_size
        bullet_run.font.bold = True
        _add_formatted_runs(p, clean_item, font_name, body_size, primary_rgb, accent_color=accent_rgb)

    if has_image:
        img_info = slide_data.images[0]
        _add_picture_fit(slide, img_info.resolved_path, Inches(7.4), Inches(2.2), Inches(4.9), Inches(4.3))
        if img_info.alt:
            cap_box = slide.shapes.add_textbox(Inches(7.4), Inches(6.6), Inches(4.9), Inches(0.4))
            cp = cap_box.text_frame.paragraphs[0]
            cp.alignment = PP_ALIGN.CENTER
            _add_formatted_runs(cp, img_info.alt, font_name, Pt(11), secondary_rgb, accent_color=accent_rgb)


def _has_audio(slide_data: SlideData) -> bool:
    return bool(slide_data.audio_path and slide_data.audio_path.exists())


_HEADER_TITLE_PT = 32
_MIN_HEADER_TITLE_PT = 20
_HEADER_TITLE_H = Inches(1.1)
_AUDIO_WIDGET_LEFT = Inches(10.2)


def _add_header(slide, category: str, slide_data: SlideData, accent_rgb, primary_rgb, theme: ThemeConfig):
    font_name = _get_font_name(theme)
    title_font = _get_title_font(theme)
    arch = theme.archetype
    title = slide_data.title

    # The title stops short of the audio widget; it shrinks from 32 pt until it fits without a mid-word break.
    title_left = Inches(1.0)
    title_w = (_AUDIO_WIDGET_LEFT - Inches(0.15) - title_left) if _has_audio(slide_data) else Inches(11.0)
    width_pt = (title_w - _BOX_INSET_X) / 12700
    avail_pt = (_HEADER_TITLE_H - _BOX_INSET_Y) / 12700
    title_pt = _HEADER_TITLE_PT
    while True:
        wrapped = wrap_markdown(title, title_font, title_pt, width_pt, bold=True)
        fits = not wrapped.broke_word and wrapped.lines * line_height_pt(title_font, title_pt) <= avail_pt
        if fits or title_pt <= _MIN_HEADER_TITLE_PT:
            break
        title_pt -= 1
    if not fits:
        print(f"[PPTX] Warning: slide {slide_data.index}: title does not fit even at {title_pt} pt; shorten it.")
    title_size = Pt(title_pt)

    # Category lead
    cat_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.5), Inches(8.0), Inches(0.35))
    cp = cat_box.text_frame.paragraphs[0]
    prefix = "root@slides:~# " if arch == ArchetypeType.TERMINAL else ""
    cp.text = f"{prefix}{category}"
    cp.font.name = font_name
    cp.font.size = Pt(13)
    cp.font.bold = True
    cp.font.color.rgb = accent_rgb

    # Slide Title
    title_box = slide.shapes.add_textbox(title_left, Inches(0.85), title_w, _HEADER_TITLE_H)
    tf = title_box.text_frame
    tf.word_wrap = True
    tp = tf.paragraphs[0]
    _add_formatted_runs(tp, title, title_font, title_size, primary_rgb, is_title=True, accent_color=accent_rgb)

    # Editorial hairline header separator
    if arch == ArchetypeType.EDITORIAL:
        hdr_rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.95), Inches(11.333), Inches(0.015))
        hdr_rule.fill.solid()
        hdr_rule.fill.fore_color.rgb = RGBColor(*theme.hex_to_rgb(theme.colors.get('card_border', '#EADBCE')))
        hdr_rule.line.fill.background()


